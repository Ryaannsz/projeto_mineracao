from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from models import VendaConcorrenteOrigem

_MESES = {
    "jan": 1,
    "fev": 2,
    "mar": 3,
    "abr": 4,
    "mai": 5,
    "jun": 6,
    "jul": 7,
    "ago": 8,
    "set": 9,
    "out": 10,
    "nov": 11,
    "dez": 12,
}


class RepositorioPlanilhaConcorrente:
    """Adaptador da planilha Excel com as vendas mensais da concorrente."""

    def __init__(self, caminho_planilha: str | Path) -> None:
        self._caminho_planilha = Path(caminho_planilha)

    def extrair(self) -> tuple[VendaConcorrenteOrigem, ...]:
        pasta_trabalho = load_workbook(self._caminho_planilha, read_only=True, data_only=True)
        planilha = pasta_trabalho.active

        vendas: list[VendaConcorrenteOrigem] = []
        linhas = planilha.iter_rows(min_row=2, values_only=True)
        for ano, mes, valor in linhas:
            if ano is None and mes is None and valor is None:
                continue
            vendas.append(
                VendaConcorrenteOrigem(
                    ano=int(ano),
                    mes=_para_mes(mes),
                    valor=Decimal(str(valor)),
                )
            )

        pasta_trabalho.close()
        return tuple(vendas)


def _para_mes(valor: object) -> int:
    if isinstance(valor, int):
        return valor
    chave = str(valor).strip().lower()[:3]
    mes = _MESES.get(chave)
    if mes is None:
        raise ValueError(f"Mes nao reconhecido na planilha da concorrente: {valor!r}")
    return mes
